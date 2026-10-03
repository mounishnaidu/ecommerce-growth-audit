"""Generator for the E-commerce Growth Audit project plan docx."""
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

OUT = "/home/hatch/workspace/your_files/ecommerce-growth-audit/Ecommerce Growth Audit - Project Plan.docx"

doc = Document()
doc.core_properties.title = "E-commerce Growth Audit - Project Plan"
doc.core_properties.author = "Shadow"

style = doc.styles["Normal"]
style.font.name = "Calibri"
style.font.size = Pt(11)

def h1(text):
    p = doc.add_heading(text, level=1)
    return p

def h2(text):
    doc.add_heading(text, level=2)

def para(text, bold=False):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold = bold
    return p

def bullets(items):
    for it in items:
        doc.add_paragraph(it, style="List Bullet")

def numbered(items):
    for it in items:
        doc.add_paragraph(it, style="List Number")

# Title block
t = doc.add_paragraph()
r = t.add_run("E-commerce Growth Audit")
r.bold = True
r.font.size = Pt(22)
t.alignment = WD_ALIGN_PARAGRAPH.LEFT
sub = doc.add_paragraph()
r = sub.add_run("Portfolio Project Plan  |  Started October 2026")
r.font.size = Pt(12)
r.font.color.rgb = RGBColor(0x59, 0x59, 0x59)

h1("Project Overview")
para("An end-to-end growth marketing analytics project built on a real e-commerce dataset. "
     "It applies every skill from the learning track — SQL, marketing analytics, statistics, "
     "Excel, Python, BI, GA4/Amplitude thinking, and MarTech — and ends with a public portfolio "
     "piece: a GitHub repo, an interactive dashboard, and a written growth audit.")
bullets([
    "Dataset: Olist Brazilian Marketplace (~100,000 orders, 2016-2018) — customers, orders, payments, reviews, products, sellers, plus a marketing-qualified-leads funnel.",
    "Timeline: ~12 days at 2-3 hours per day.",
    "Deliverables: clean database, SQL analysis, churn model, Tableau Public dashboard, 2-page growth audit memo, GitHub repo, resume bullet.",
])

h1("Roadmap")
phases = [
    ("Phase 1 - Data Foundations (Days 1-2)",
     ["Download the Olist dataset from Kaggle.",
      "Load CSVs into SQLite and pandas.",
      "Write a data dictionary; run quality checks (missing values, duplicates, date ranges)."],
     "Clean SQLite database + exploratory analysis notebook."),
    ("Phase 2 - Funnel & Channel Analytics (Days 3-5)",
     ["Marketing funnel: MQL to SQL to closed-order conversion rates; find where it leaks.",
      "Revenue by product category, customer state, and payment type.",
      "Cohort analysis: repeat-purchase rate by acquisition month."],
     "SQL query folder + written findings."),
    ("Phase 3 - Customer Economics (Days 6-8)",
     ["RFM segmentation; estimate LTV per segment.",
      "Define churn (e.g., no order in 6 months); build a churn classifier (logistic regression, then random forest); report feature importances."],
     "Modeling notebook with evaluation metrics."),
    ("Phase 4 - Dashboard (Days 9-10)",
     ["Build a Tableau Public dashboard: KPI overview, funnel, cohort heatmap, churn-risk segment."],
     "Public dashboard link."),
    ("Phase 5 - Report & Ship (Days 11-12)",
     ["Write a 2-page growth audit memo: 3 findings + 3 recommendations with expected impact.",
      "Publish the GitHub repo with a README linking notebooks, dashboard, and memo.",
      "Add the resume bullet."],
     "Shipped portfolio project."),
]
for title, steps, deliverable in phases:
    h2(title)
    numbered(steps)
    para("Deliverable: " + deliverable, bold=False)

h2("Stretch Goals (Optional)")
bullets([
    "Design an A/B test for checkout conversion (hypothesis, metric, sample size).",
    "Write a GA4 event spec for the store (events, parameters, conversions).",
])

h1("How We Work")
bullets([
    "Mounish commands, Shadow executes: every step starts with a specific instruction from Mounish.",
    "Shadow explains the why behind each step, not just the what.",
    "Commands must be specific — vague orders get sent back for sharpening.",
    "Before moving to the next phase, Mounish explains the decisions back in his own words (interview-ready).",
    "One phase at a time; no advancing on broken prior work.",
])

h1("Progress Log")
table = doc.add_table(rows=1, cols=3)
table.style = "Light Grid Accent 1"
hdr = table.rows[0].cells
hdr[0].text = "Phase"
hdr[1].text = "Status"
hdr[2].text = "Notes"
progress = [
    ("Phase 1", "Complete",
     "Olist + marketing funnel datasets downloaded from Kaggle (11 CSVs). Loaded into SQLite "
     "(data/olist.db, 11 tables, indexes on join keys). Quality-checked and cleaned: 261,831 "
     "duplicate geolocation rows removed; 610 missing product categories filled 'unknown'; "
     "informative missingness (delivery dates, review text) preserved. QA log at data/qa_report.txt. "
     "Data dictionary built: Olist Data Dictionary.xlsx (in Drive)."),
    ("Phase 2", "Not started", ""),
    ("Phase 3", "Not started", ""),
    ("Phase 4", "Not started", ""),
    ("Phase 5", "Not started", ""),
]
for title, status, notes in progress:
    row = table.add_row().cells
    row[0].text = title
    row[1].text = status
    row[2].text = notes

doc.save(OUT)
print("saved", OUT)
